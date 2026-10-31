import csv
import io
from datetime import date

import jdatetime
import pytest
from openpyxl import load_workbook

BASE = "/api/qaem"
RANGE = {"from": "2026-09-01", "to": "2026-10-31"}
SEED = [("2026-09-30", 1, 10), ("2026-09-30", 2, 4), ("2026-10-01", 1, 5), ("2026-10-22", 7, 0)]
TITLE = {"services": "سرویس‌های ثبت‌شده", "totals": "مجموع سرویس‌ها"}
HEAD = {"services": ["تاریخ", "تاریخ میلادی", "شماره کامیون", "تعداد سرویس"],
        "totals": ["تاریخ", "تاریخ میلادی", "جمع روز", "جمع ماه تا این روز"]}


def seed(client):
    for d, t, c in SEED:
        assert client.post(f"{BASE}/services", json={"service_day": d, "truck_id": t, "truck_service_count": c}).status_code == 201


def solar(d):
    j = jdatetime.date.fromgregorian(date=date.fromisoformat(d))
    return f"{j.year:04d}/{j.month:02d}/{j.day:02d}"


def expected(client, table):
    """The rows the screen shows (the JSON API), in export column order."""
    keys = ("truck_id", "truck_service_count") if table == "services" else ("day_total", "month_total")
    data = client.get(f"{BASE}/{table}", params=RANGE).json()
    return [[solar(r["service_day"]), r["service_day"], *(r[k] for k in keys)] for r in data]


@pytest.mark.parametrize("table", ["services", "totals"])
def test_xlsx_matches_screen_data(client, table):
    seed(client)
    r = client.get(f"{BASE}/export/{table}", params=RANGE)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert r.headers["content-disposition"] == f'attachment; filename="qaem-{table}_2026-09-01_2026-10-31.xlsx"'
    assert r.content[:2] == b"PK"  # a real zip/xlsx container
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    assert rows[0] == [TITLE[table], None, None, None]  # table name as the heading
    assert rows[1] == HEAD[table]
    got = [[a, b.date().isoformat(), c, d] for a, b, c, d in rows[2:]]
    assert got == expected(client, table) and len(got) > 0
    assert ws.sheet_view.rightToLeft is True
    assert ws["A1"].font.b and ws["A3"].data_type == "s" and ws["B3"].is_date and ws["C3"].data_type == "n"  # text / date / number cells


@pytest.mark.parametrize("table", ["services", "totals"])
def test_csv_matches_screen_data_and_has_bom_for_excel(client, table):
    seed(client)
    r = client.get(f"{BASE}/export/{table}", params={**RANGE, "format": "csv"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "text/csv; charset=utf-8"
    assert r.headers["content-disposition"].endswith('.csv"')
    assert r.content.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM so Excel shows Farsi correctly
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == HEAD[table]
    assert [[a, b, int(c), int(d)] for a, b, c, d in rows[1:]] == expected(client, table)


def test_known_solar_value(client):
    seed(client)
    rows = list(csv.reader(io.StringIO(client.get(f"{BASE}/export/services", params={**RANGE, "format": "csv"}).text.lstrip("\ufeff"))))
    assert rows[1][:2] == ["1405/07/08", "2026-09-30"]  # 2026-09-30 is 8 Mehr 1405


def test_export_respects_range_and_empty_range_gives_header_only(client):
    seed(client)
    r = client.get(f"{BASE}/export/services", params={"from": "2026-10-01", "to": "2026-10-01", "format": "csv"})
    assert len(list(csv.reader(io.StringIO(r.text.lstrip("\ufeff"))))) == 2
    r = client.get(f"{BASE}/export/totals", params={"from": "2020-01-01", "to": "2020-01-02"})
    assert [list(x) for x in load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True)] == [[TITLE["totals"], None, None, None], HEAD["totals"]]


def test_zero_count_and_large_values_survive(client):
    client.post(f"{BASE}/services", json={"service_day": "2026-10-01", "truck_id": 2147483647, "truck_service_count": 2147483647})
    client.post(f"{BASE}/services", json={"service_day": "2026-10-01", "truck_id": 1, "truck_service_count": 0})
    ws = load_workbook(io.BytesIO(client.get(f"{BASE}/export/services", params=RANGE).content)).active
    assert [row[2:] for row in ws.iter_rows(min_row=3, values_only=True)] == [(1, 0), (2147483647, 2147483647)]


@pytest.mark.parametrize("params", [
    {"format": "pdf"}, {"format": "XLSX"}, {"from": "2026-10-02", "to": "2026-10-01"},
    {"from": "2020-01-01", "to": "2026-01-01"}, {"from": "x"}, {},
])
def test_export_rejects_bad_requests(client, params):
    p = {**RANGE, **params} if "from" not in params and params else params
    assert client.get(f"{BASE}/export/services", params=p).status_code == 422


def test_unknown_table_is_rejected(client):
    assert client.get(f"{BASE}/export/users", params=RANGE).status_code == 422


def test_default_format_is_xlsx(client):
    assert client.get(f"{BASE}/export/services", params=RANGE).content[:2] == b"PK"