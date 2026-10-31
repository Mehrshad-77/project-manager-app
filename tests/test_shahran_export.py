import csv
import io
from datetime import date, time

import jdatetime
import pytest
import psycopg
from openpyxl import load_workbook

BASE = "/api/shahran"
ABAN = {"from": "2026-10-23", "to": "2026-11-21"}          # 1 Aban 1405 .. 30 Aban 1405
TITLE = {"time-sheet": "کارکرد روزانه", "aggregated-work": "کارکرد تجمیعی ماهانه"}
HEAD = {
    "time-sheet": ["تاریخ", "تاریخ میلادی", "کارگر", "ساعت ورود", "ساعت خروج", "اضافه‌کاری", "تعطیلات", "مرخصی", "توضیحات"],
    "aggregated-work": ["تاریخ", "تاریخ میلادی", "مجموع کارگران", "کارگر ساده", "بنا", "توضیحات"],
}


@pytest.fixture(autouse=True)
def clean_shahran(database):
    with psycopg.connect(pytest.admin_url, autocommit=True) as conn:
        conn.execute("TRUNCATE shahran.daily_time_sheet, shahran.daily_notes, shahran.workers RESTART IDENTITY CASCADE")


def worker(client, name, role, salary=1000):
    r = client.post(f"{BASE}/workers", json={"full_name": name, "role": role, "daily_salary": salary})
    assert r.status_code == 201, r.text
    return r.json()["worker_id"]


def sheet(client, day, wid, **kw):
    r = client.post(f"{BASE}/time-sheet", json={"work_day": day, "worker_id": wid, **kw})
    assert r.status_code == 201, r.text


def seed(client):
    ali, hasan = worker(client, "Ali", "simple"), worker(client, "Hasan", "builder")
    sheet(client, "2026-10-23", ali, arrival_time="07:00", exit_time="16:00", overtime=2.5, explanation="late")
    sheet(client, "2026-10-23", hasan, is_off_day=True)
    sheet(client, "2026-10-24", hasan, arrival_time="08:00", exit_time="17:30", is_holiday=True)
    assert client.put(f"{BASE}/aggregated-work/2026-10-25", json={"explanation": "rain"}).status_code == 200


def xlsx_rows(r):
    return [list(x) for x in load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True)]


def csv_rows(r):
    return list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))


def get(client, table, fmt="xlsx", **params):
    return client.get(f"{BASE}/export/{table}", params={**ABAN, "format": fmt, **params})


# ---------- the time sheet (کارکرد روزانه): hand-written expected rows ----------
def test_time_sheet_xlsx_exact_rows(client):
    seed(client)
    r = get(client, "time-sheet")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert r.headers["content-disposition"] == 'attachment; filename="shahran-time-sheet_2026-10-23_2026-11-21.xlsx"'
    rows = xlsx_rows(r)
    assert rows[0] == ["کارکرد روزانه"] + [None] * 8                      # table name as the heading
    assert rows[1] == HEAD["time-sheet"]
    d = lambda y, m, dd: __import__("datetime").datetime(y, m, dd)
    assert rows[2:] == [
        ["1405/08/01", d(2026, 10, 23), "Ali", time(7, 0), time(16, 0), 2.5, None, None, "late"],
        ["1405/08/01", d(2026, 10, 23), "Hasan", None, None, 0, None, "✓", None],
        ["1405/08/02", d(2026, 10, 24), "Hasan", time(8, 0), time(17, 30), 0, "✓", None, None],
    ]


def test_time_sheet_xlsx_cell_types_and_layout(client):
    seed(client)
    ws = load_workbook(io.BytesIO(get(client, "time-sheet").content)).active
    assert ws.title == "time-sheet" and ws.sheet_view.rightToLeft is True and ws["A1"].font.b
    assert ws["A3"].data_type == "s" and ws["B3"].is_date and ws["D3"].number_format == "hh:mm" and ws["F3"].data_type == "n"


def test_time_sheet_csv_exact_rows(client):
    seed(client)
    r = get(client, "time-sheet", "csv")
    assert r.content.startswith(b"\xef\xbb\xbf") and r.headers["content-type"] == "text/csv; charset=utf-8"
    assert csv_rows(r) == [
        HEAD["time-sheet"],
        ["1405/08/01", "2026-10-23", "Ali", "07:00", "16:00", "2.5", "", "", "late"],
        ["1405/08/01", "2026-10-23", "Hasan", "", "", "0.0", "", "✓", ""],
        ["1405/08/02", "2026-10-24", "Hasan", "08:00", "17:30", "0.0", "✓", "", ""],
    ]


# ---------- the aggregated sheet (کارکرد تجمیعی ماهانه) ----------
def test_aggregated_xlsx_has_every_day_and_right_counts(client):
    seed(client)
    rows = xlsx_rows(get(client, "aggregated-work"))
    assert rows[0] == ["کارکرد تجمیعی ماهانه"] + [None] * 5
    assert rows[1] == HEAD["aggregated-work"]
    body = rows[2:]
    assert len(body) == 30                                                 # every day of Aban, like the screen
    by_day = {r[0]: r[2:] for r in body}
    assert by_day["1405/08/01"] == [1, 1, 0, None]                         # Ali worked; Hasan's off-day row not counted
    assert by_day["1405/08/02"] == [1, 0, 1, None]                         # Hasan (builder) worked on a holiday
    assert by_day["1405/08/03"] == [0, 0, 0, "rain"]                       # 2026-10-25 note only
    assert by_day["1405/08/30"] == [0, 0, 0, None]
    assert body[0][1].date().isoformat() == "2026-10-23" and body[-1][1].date().isoformat() == "2026-11-21"


def test_aggregated_csv(client):
    seed(client)
    rows = csv_rows(get(client, "aggregated-work", "csv"))
    assert rows[0] == HEAD["aggregated-work"] and len(rows) == 31
    assert rows[1] == ["1405/08/01", "2026-10-23", "1", "1", "0", ""]
    assert rows[3] == ["1405/08/03", "2026-10-25", "0", "0", "0", "rain"]


# ---------- the export is exactly what the screen's API returns ----------
def test_matches_api_for_both_tables(client):
    seed(client)
    ts = client.get(f"{BASE}/time-sheet", params=ABAN).json()
    got = xlsx_rows(get(client, "time-sheet"))[2:]
    assert [(r[2], r[3] and r[3].strftime("%H:%M:%S"), r[4] and r[4].strftime("%H:%M:%S")) for r in got] == \
           [(x["full_name"], x["arrival_time"], x["exit_time"]) for x in ts]
    agg = client.get(f"{BASE}/aggregated-work", params=ABAN).json()
    got = xlsx_rows(get(client, "aggregated-work"))[2:]
    assert [(r[1].date().isoformat(), *r[2:5], r[5]) for r in got] == \
           [(x["work_day"], x["total_workers"], x["simple_workers"], x["builders"], x["explanation"]) for x in agg]


def test_solar_dates_match_jdatetime_for_every_exported_day(client):
    got = xlsx_rows(get(client, "aggregated-work"))[2:]
    for solar, g, *_ in got:
        j = jdatetime.date.fromgregorian(date=g.date())
        assert solar == f"{j.year:04d}/{j.month:02d}/{j.day:02d}"


def test_empty_database(client):
    assert xlsx_rows(get(client, "time-sheet")) == [["کارکرد روزانه"] + [None] * 8, HEAD["time-sheet"]]
    assert csv_rows(get(client, "time-sheet", "csv")) == [HEAD["time-sheet"]]
    assert len(xlsx_rows(get(client, "aggregated-work"))) == 32           # title + header + 30 zero days


def test_range_limits_rows(client):
    seed(client)
    r = client.get(f"{BASE}/export/time-sheet", params={"from": "2026-10-24", "to": "2026-10-24", "format": "csv"})
    assert [x[2] for x in csv_rows(r)[1:]] == ["Hasan"]


# ---------- typed text must never turn into a spreadsheet formula ----------
def test_formula_like_text_stays_text(client):
    evil = worker(client, "=1+1", "simple")
    other = worker(client, "-x", "builder")
    sheet(client, "2026-10-23", evil, explanation='=HYPERLINK("http://evil","click")')
    sheet(client, "2026-10-23", other, explanation="+cmd|calc")
    assert client.put(f"{BASE}/aggregated-work/2026-10-23", json={"explanation": "@SUM(1)"}).status_code == 200
    ws = load_workbook(io.BytesIO(get(client, "time-sheet").content)).active
    cells = {c.value: c.data_type for row in ws.iter_rows(min_row=3) for c in row if isinstance(c.value, str)}
    assert cells['=HYPERLINK("http://evil","click")'] == "s" and cells["=1+1"] == "s" and cells["+cmd|calc"] == "s"
    rows = csv_rows(get(client, "time-sheet", "csv"))[1:]
    assert sorted((r[2], r[8]) for r in rows) == [("'-x", "'+cmd|calc"), ("'=1+1", "'=HYPERLINK(\"http://evil\",\"click\")")]
    assert csv_rows(get(client, "aggregated-work", "csv"))[1][5] == "'@SUM(1)"


# ---------- bad requests ----------
@pytest.mark.parametrize("table", ["time-sheet", "aggregated-work"])
@pytest.mark.parametrize("params", [
    {"format": "pdf"}, {"format": "XLSX"}, {"from": "2026-10-24", "to": "2026-10-23"},
    {"from": "2020-01-01", "to": "2026-01-01"}, {"from": "x"}, {"to": None},
])
def test_bad_requests_are_422(client, table, params):
    p = {**ABAN, "format": "xlsx", **params}
    p = {k: v for k, v in p.items() if v is not None}
    assert client.get(f"{BASE}/export/{table}", params=p).status_code == 422


def test_unknown_table_and_default_format(client):
    assert client.get(f"{BASE}/export/workers", params=ABAN).status_code == 422
    assert client.get(f"{BASE}/export/monthly-salary", params=ABAN).status_code == 422
    assert client.get(f"{BASE}/export/time-sheet", params=ABAN).content[:2] == b"PK"