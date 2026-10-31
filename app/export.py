"""Shared by every project: turn rows into a downloadable Excel (.xlsx) or CSV response."""
import csv
import io
from datetime import date, time

import jdatetime
from fastapi import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DAY_HEADERS = ("تاریخ", "تاریخ میلادی")  # Solar date (text, as on screen), Gregorian date (a real date cell)


def day_cells(d: date) -> list:
    j = jdatetime.date.fromgregorian(date=d)
    return [f"{j.year:04d}/{j.month:02d}/{j.day:02d}", d]


def _csv_safe(v):
    """Excel runs text that starts with = + - @ as a formula; a leading quote keeps it text."""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def export_file(project, table, fmt, rng, title, headers, rows):
    """rows: tuples whose first item is the day (a date); the rest follow `headers`.
    The xlsx has `title` (the table's name) as a bold heading in row 1; the CSV is a plain table."""
    headers = [*DAY_HEADERS, *headers]
    rows = [[*day_cells(r[0]), *r[1:]] for r in rows]
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = table
        ws.sheet_view.rightToLeft = True
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=3):
            for c in row:
                if c.data_type == "f":  # typed text such as "=1+1" must stay text
                    c.data_type = "s"
                elif isinstance(c.value, time):
                    c.number_format = "hh:mm"
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 20
        out = io.BytesIO()
        wb.save(out)
        content, media = out.getvalue(), XLSX_TYPE
    else:
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(headers)
        for r in rows:
            w.writerow([v.strftime("%H:%M") if isinstance(v, time) else _csv_safe(v) for v in r])
        content, media = buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"  # BOM: Excel reads Farsi correctly
    name = f"{project}-{table}_{rng[0]}_{rng[1]}.{fmt}"
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="{name}"'})